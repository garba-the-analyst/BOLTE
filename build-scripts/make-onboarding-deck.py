#!/usr/bin/env python3
"""Build BOLTE onboarding deck — 9 slides, 16:9, dark theme."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BG = RGBColor(0, 0, 0)
BLUE = RGBColor(0x2E, 0x86, 0xDE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xB8, 0xC2, 0xCC)
CARD = RGBColor(0x12, 0x14, 0x1A)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

def add_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

def textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))

def para(tf, text, size=20, bold=False, color=WHITE, align=PP_ALIGN.LEFT, space_after=Pt(4)):
    p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    p.space_after = space_after
    p.font.name = "Calibri"
    return p

def title_slide(slide, kicker, title, subtitle, notes):
    add_bg(slide)
    tb = textbox(slide, 0.8, 0.5, 11.7, 1.0)
    para(tb.text_frame, kicker, size=16, bold=True, color=BLUE, align=PP_ALIGN.LEFT)
    tb2 = textbox(slide, 0.8, 1.5, 11.7, 2.2)
    tf = tb2.text_frame; tf.word_wrap = True
    para(tf, title, size=44, bold=True, color=WHITE)
    tb3 = textbox(slide, 0.8, 3.9, 11.7, 2.2)
    tf3 = tb3.text_frame; tf3.word_wrap = True
    for line in subtitle:
        para(tf3, line, size=18, color=MUTED)
    slide.notes_slide.placeholders[1].text = notes

def content_slide(title, bullets, notes, badge=None):
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    tb = textbox(slide, 0.8, 0.3, 11.7, 1.1)
    para(tb.text_frame, title, size=32, bold=True, color=WHITE)
    # blue rule
    shape = slide.shapes.add_shape(1, Inches(0.8), Inches(1.35), Inches(2.0), Pt(4))
    shape.fill.solid(); shape.fill.fore_color.rgb = BLUE; shape.line.fill.background()
    tb2 = textbox(slide, 0.8, 1.7, 11.7, 4.8)
    tf = tb2.text_frame; tf.word_wrap = True
    for b in bullets:
        p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p.text = b; p.level = 0
        p.font.size = Pt(20); p.font.color.rgb = WHITE; p.font.name = "Calibri"
        p.space_after = Pt(8)
    if badge:
        tb3 = textbox(slide, 9.8, 0.35, 2.7, 0.8)
        para(tb3.text_frame, badge, size=14, bold=True, color=BLUE, align=PP_ALIGN.RIGHT)
    slide.notes_slide.placeholders[1].text = notes
    return slide

# 1 Title
s = prs.slides.add_slide(blank)
title_slide(s, "BOLTE  •  CORPORATE DIRECTION", "BOLTE",
    ["Indigenous Technology for Nigeria — and ultimately Africa.",
     "First Onboarding  •  Virtual  •  60 minutes  •  Co-chaired"],
    "0:00-0:05 welcome. One-mic, 60-sec intros coming, [TBD] acceptable.")
# 2 Identity
content_slide("Identity — Vision / Mission / Aim", [
    "VISION: a technologically independent Nigeria, ultimately Africa.",
    "MISSION: research, engineer, develop, deploy indigenous tech — defense & security first.",
    "AIM: self-reliance through research, engineering, invention — foundation for Africa.",
    "Bigger cycle: Nigeria → African Capability → African Technological Independence.",
], "0:15-0:25 story. 2 min vision, 5 min goals buckets, 3 min model. No debate.")
# 3 Goals
content_slide("7 Goals at a glance", [
    "G1 Defense & security technology (primary focus).",
    "G2 Backbone  •  G3 Ideas→inventions  •  G4 Talent community.",
    "G5 Strategic sectors  •  G6 Research ecosystem  •  G7 Build FOR Nigeria.",
    "Rule: if tech solves an important Nigerian problem, explore it.",
], "Buckets only, details in docs/01-canon.")
# 4 Model
content_slide("The BOLTE Model", [
    "PROBLEM → IDEA → RESEARCH → ENGINEERING → PROTOTYPE →",
    "TESTING → IMPROVEMENT → DEPLOYMENT → IMPACT.",
    "Ecosystem bridge: Idea → Research → Prototype → Testing → Product → Deployment.",
    "“All problems are solved when you put things together.”",
], "3 min. Multidisciplinary emphasis.")
# 5 How we work
content_slide("How we work", [
    "2 founders + 5 seats: Treasurer, Product/Ops, 3 builders (assignments async).",
    "Cadence: weekly build review + monthly strategy. Everything documented.",
    "Conduct: original work, no leaks, declare conflicts, safety first.",
    "Money (detail next session): 50% treasury / 20% founders / 30% members.",
], "0:25-0:35. No seat assignments today — capture preferences only.")
# 6 Flagship
content_slide("Flagship: AeroPulse-NG", [
    "Air-gapped offline-first surveillance: ~$150 SDR + PC → ATC/tactical workstation.",
    "ADS-B/Mode S + ACARS → EKF tracking → STCA conflicts → weather fusion → defence overlays.",
    "Single decode path (sim == live). Deterministic, replayable, locally maintainable.",
], "3-min demo. Honest status, no hype.", badge="PROTOTYPE • TRL 4–5")
# 7 NDAIE
content_slide("Immediate focus: NDAIE 2026", [
    "National Innovation Pitch — Flight Planning & Air Traffic Management.",
    "Due 26 Sept 23:59 WAT: form + 300-word summary + 10-slide PDF.",
    "Only 2–5 registered students (16–30, enrolled) may present.",
    "Internal freeze 19 Sept • submit 25 Sept • video 14 Oct • semi 28 Oct • finals 27 Nov Lagos.",
], "0:35-0:48. Capture eligibility Y/N in chat today.")
# 8 Next
content_slide("What happens next", [
    "Homework (48h): questionnaire minimum — date, seat preference, eligibility, deck comments.",
    "Next session: lock sub-team + slide owners + travel plan.",
    "Nothing to sign today. MOU/NDA/IP + equity finalized async after Session 2.",
    "Access: confirm comms channel + file sharing before close.",
], "0:48-0:56.")
# 9 Close
content_slide("Close — put things together", [
    "“All problems are solved when you put things together.”",
    "Q&A (remaining minutes).",
    "Action log readout: Decision | Owner | Date — verbal consent to close.",
    "Thank you — build for Nigeria.",
], "0:56-1:00. End on time.")

out = "/home/garba-the-analyst/Desktop/Projects/BOLTE/BOLTE-Onboarding-Deck.pptx"
prs.save(out)
print(f"saved {out} — {len(prs.slides)} slides")
